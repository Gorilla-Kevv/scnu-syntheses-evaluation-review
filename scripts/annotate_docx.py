# -*- coding: utf-8 -*-
"""annotate_docx.py —— 通用 Word 原生批注生成器。

在**副本**上生成原生 Word 批注，原件绝不改动。副本落在原文件同级目录。

依赖：python-docx >= 1.2.0（1.1.x 不支持 add_comment）

批注定义文件（JSON）格式::

    {
      "src": "F:/path/xxx.docx",
      "dst": "F:/path/xxx（审核批注版）.docx",       // 可省略，默认同级目录
      "author": "综测审核组",
      "initials": "SJ",
      "overview": {                                  // 可选：总述批注
        "anchor": ["hval", "姓名"],
        "text": "【审核总述】……（并列出不可读材料清单）",
        "fallback": ["lbl", 0, "姓名"]
      },
      "notes": [
        {"anchor": ["val", 14, "劳动观念加分"], "text": "【问题1·分值应更正】……"},
        {"anchor": ["kw", "作为队员参加寒假社会实践专项调研比赛"], "text": "【问题2·……】"},
        {"anchor": ["hval", "年级"], "text": "【问题3·……】"}
      ]
    }

锚点类型：
    ["hval", 字段]         表头：字段的值单元格
    ["hlbl", 字段]         表头：字段的标签单元格
    ["val", 行号, 字段]     第 N 行的值单元格（为空时自动回退到标签单元格）
    ["lbl", 行号, 字段]     第 N 行的标签单元格
    ["kw", 关键词]          含关键词的段落（忽略空白回退匹配）

用法：
    python annotate_docx.py notes.json
    python annotate_docx.py notes.json --folder F:/path/学生文件夹   # 附带扫描不可读材料清单
"""
import argparse
import json
import os
import shutil
import sys
import zipfile

try:
    from docx import Document
except ImportError:
    print("需要 python-docx >= 1.2.0：pip install -U python-docx")
    sys.exit(1)


# ---------------- 基础工具 ----------------
def row_cells(row):
    """按 tc 去重返回单元格（合并单元格在 row.cells 中会重复）。"""
    seen, out = set(), []
    for c in row.cells:
        k = id(c._tc)
        if k in seen:
            continue
        seen.add(k)
        out.append(c)
    return out


def find_label(doc, row_idx, label):
    cells = row_cells(doc.tables[0].rows[row_idx])
    lab = label.replace(" ", "")
    for i, c in enumerate(cells):
        if c.text.strip().replace(" ", "") == lab:
            return c, (cells[i + 1] if i + 1 < len(cells) else None)
    return None, None


def all_paragraphs(doc):
    """遍历正文 + 全部表格内的段落。"""
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


def find_kw_para(doc, kw):
    for p in all_paragraphs(doc):
        if kw in p.text:
            return p
    # 忽略空白回退（部分表内文字被逐字加空格）
    k = kw.replace(" ", "")
    for p in all_paragraphs(doc):
        if k in p.text.replace(" ", ""):
            return p
    return None


def run_of(cell):
    for p in cell.paragraphs:
        if p.runs:
            return p.runs
    return [cell.paragraphs[0].add_run("")]


def run_of_para(p):
    if p.runs:
        return p.runs
    return [p.add_run("")]


def resolve(doc, anchor):
    """返回 (runs, 位置描述) 或 (None, 失败原因)。"""
    kind = anchor[0]
    if kind in ("hval", "hlbl"):
        lab, val = find_label(doc, 0, anchor[1])
        cell = val if (kind == "hval" and val is not None) else lab
        if cell is None:
            return None, "表头字段未找到：%s" % anchor[1]
        return run_of(cell), "表头·%s" % anchor[1]
    if kind in ("val", "lbl"):
        lab, val = find_label(doc, anchor[1], anchor[2])
        if lab is None:
            return None, "单元格未找到：R%d %s" % (anchor[1], anchor[2])
        if kind == "val" and val is not None and val.text.strip():
            return run_of(val), "R%d·%s" % (anchor[1], anchor[2])
        return run_of(lab), "R%d·%s(标签)" % (anchor[1], anchor[2])
    if kind == "kw":
        p = find_kw_para(doc, anchor[1])
        if p is None:
            return None, "未找到文字：%s" % anchor[1]
        return run_of_para(p), "文字[%s]" % anchor[1]
    return None, "未知锚点类型：%r" % (anchor,)


# ---------------- 不可读材料清单 ----------------
def unreadable_list(folder):
    """扫描文件夹，返回写不出文字的图片/扫描PDF（用于总述批注说明）。"""
    res = []
    try:
        import pymupdf
    except ImportError:
        pymupdf = None
    for dp, dn, fn in os.walk(folder):
        for f in sorted(fn):
            if "审核批注版" in f or "审核修订版" in f:
                continue
            p = os.path.join(dp, f)
            low = f.lower()
            if low.endswith((".jpg", ".jpeg", ".png", ".bmp", ".gif")):
                res.append((f, "图片文件（无文字层），无法直接提取文字核对"))
            elif low.endswith(".pdf"):
                if pymupdf is None:
                    continue
                try:
                    d = pymupdf.open(p)
                    t = "".join(pg.get_text() for pg in d)
                    if len(t.strip()) < 20:
                        res.append((f, "扫描版PDF（无文字层），无法直接提取文字核对"))
                except Exception as e:
                    res.append((f, "PDF无法解析：%s" % e))
    return res


# ---------------- 主流程 ----------------
def annotate(spec, folder=None):
    src = spec["src"]
    if not os.path.exists(src):
        print("!! 原件不存在：%s" % src)
        return 1

    dst = spec.get("dst") or os.path.join(
        os.path.dirname(src), "%s（审核批注版）.docx" % os.path.splitext(os.path.basename(src))[0])
    shutil.copy2(src, dst)          # 只改副本
    doc = Document(dst)

    author = spec.get("author", "综测审核组")
    initials = spec.get("initials", "SJ")
    logs, fails = [], []

    # 总述
    ov = spec.get("overview")
    if ov:
        text = ov["text"]
        if folder and "unreadable" not in ov:      # 未显式给出清单时自动附上
            ul = unreadable_list(folder)
            if ul:
                text += "\n\n◆ 本文件夹内“无法读取”的佐证材料（无文字层，请本人指认或补交含姓名版本）："
                for i, (rel, why) in enumerate(ul, 1):
                    text += "\n  %d) %s —— %s" % (i, rel, why)
        r, pos = resolve(doc, ov["anchor"])
        if r is None and ov.get("fallback"):
            r, pos = resolve(doc, ov["fallback"])
        if r is None:
            fails.append("总述锚定失败：%s" % pos)
        else:
            doc.add_comment(r, text=text, author=author, initials=initials)
            logs.append(("总述", pos))

    # 逐条问题
    for n in spec.get("notes", []):
        r, pos = resolve(doc, n["anchor"])
        if r is None:
            fails.append("%s -> %s" % (pos, n["text"][:40]))
            continue
        doc.add_comment(r, text=n["text"], author=author, initials=initials)
        logs.append(("问题", pos))

    doc.save(dst)

    # 校验
    z = zipfile.ZipFile(dst)
    has_c = "word/comments.xml" in z.namelist()
    n_start = z.read("word/document.xml").decode("utf-8").count("commentRangeStart")
    print("dst:", dst)
    print("comments.xml:", has_c, " commentRangeStart:", n_start, " 批注数:", len(logs))
    for k, p in logs:
        print("   ", k, p)
    if fails:
        print("!! FAILS:")
        for f in fails:
            print("   ", f)
    if not fails and has_c and n_start == len(logs):
        print("OK: 全部批注锚定成功，原件未改动。")
    return 0 if not fails else 2


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec", help="批注定义 JSON 文件")
    ap.add_argument("--folder", default=None, help="学生文件夹（自动附不可读材料清单）")
    args = ap.parse_args()

    with open(args.spec, "r", encoding="utf-8") as f:
        spec = json.load(f)
    if isinstance(spec, list):
        rc = 0
        for s in spec:
            rc |= annotate(s, args.folder)
        return rc
    return annotate(spec, args.folder)


if __name__ == "__main__":
    sys.exit(main())
