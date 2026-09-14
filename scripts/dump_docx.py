# -*- coding: utf-8 -*-
"""dump_docx.py —— 快速 dump docx 的段落与表格结构（含合并单元格去重）。

用途：写批注锚点前，先用本脚本看清目标文档的行号 / 字段名 / 值，再写
      ("val", 行号, "字段名") 形式的锚点。

用法：
    python dump_docx.py <file.docx> [--row 14] [--maxlen 60]

输出：
    - 非空段落列表（P0, P1, ...）
    - 每张表的行结构与去重后的单元格文本
    - --row 指定时，额外打印该行各单元格的精确文本与 run 数（定位锚点用）
"""
import argparse
import sys

from docx import Document


def row_cells(row):
    """按 tc 去重后返回单元格列表（合并单元格在 row.cells 中会重复出现）。"""
    seen, out = set(), []
    for c in row.cells:
        k = id(c._tc)
        if k in seen:
            continue
        seen.add(k)
        out.append(c)
    return out


def dump_paragraphs(doc, maxlen):
    print("=== PARAGRAPHS (non-empty) ===")
    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip()
        if t:
            print("[P%d] %s" % (i, t[:maxlen]))


def dump_tables(doc, maxlen):
    print("\n=== TABLES: %d ===" % len(doc.tables))
    for ti, tb in enumerate(doc.tables):
        print("\n---- TABLE %d  rows=%d cols=%d ----" % (ti, len(tb.rows), len(tb.columns)))
        for ri, row in enumerate(tb.rows):
            cells = [c.text.strip().replace("\n", " | ")[:maxlen] for c in row_cells(row)]
            print("R%d: %s" % (ri, cells))


def dump_row(doc, ti, ri):
    print("\n=== ROW DETAIL  table=%d row=%d ===" % (ti, ri))
    row = doc.tables[ti].rows[ri]
    for idx, c in enumerate(row_cells(row)):
        runs = len(c.paragraphs[0].runs) if c.paragraphs else 0
        print("R%d C%d: text=%r  runs=%d" % (ri, idx, c.text, runs))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path", help="docx 文件路径")
    ap.add_argument("--row", type=int, default=None, help="额外详查的行号")
    ap.add_argument("--table", type=int, default=0, help="表序号（默认 0）")
    ap.add_argument("--maxlen", type=int, default=60, help="单元格/段落截断长度")
    args = ap.parse_args()

    try:
        doc = Document(args.path)
    except Exception as e:
        print("打开失败：%s" % e)
        return 1

    dump_paragraphs(doc, args.maxlen)
    dump_tables(doc, args.maxlen)
    if args.row is not None:
        dump_row(doc, args.table, args.row)

    print("\n提示：定位登记表正文，查看 P0~P3 是否含"
          "“全日制本科生综合素质测评登记表” —— 判定是否为登记表（勿以文件名判断）。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
